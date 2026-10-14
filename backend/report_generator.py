



import json
from pathlib import Path



from datetime import datetime







from docx import Document



from docx.enum.text import WD_ALIGN_PARAGRAPH



from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT



from docx.enum.section import WD_SECTION



from docx.shared import Inches, Pt, RGBColor



from docx.oxml import OxmlElement



from docx.oxml.ns import qn











# ============================================================



# IDENTIDADE VISUAL



# ============================================================







DARK_BLUE = "17365D"

GREEN = "24EBA9"

DARK_GRAY = "2D2D2D"

WHITE = "FFFFFF"

LIGHT_GREEN = "F0F8F3"

BORDER_GREEN = "D5E6DC"



TECHNICAL_TABLE_PREFIXES = (



    "LocalDateTable_",



    "DateTableTemplate_",



)











# ============================================================



# FILTROS DE APRESENTAÇÃO



# ============================================================







def is_technical_table(name: str) -> bool:



    return any(



        name.startswith(prefix)



        for prefix in TECHNICAL_TABLE_PREFIXES



    )











def is_calculated_column(column: dict) -> bool:

    """Identifica colunas calculadas a partir dos metadados extraídos do TMDL."""

    if column.get("is_calculated") is True or column.get("isCalculated") is True:

        return True



    expression = column.get("expression") or column.get("formula")

    return isinstance(expression, str) and bool(expression.strip())





def get_business_tables(model: dict) -> list:



    return [



        table



        for table in model.get("tables", [])



        if not is_technical_table(table.get("name", ""))



    ]











def get_business_relationships(model: dict) -> list:



    technical_names = {



        table.get("name", "")



        for table in model.get("tables", [])



        if is_technical_table(table.get("name", ""))



    }







    relationships = []







    for relationship in model.get("relationships", []):



        source = relationship.get("from_column", "")



        target = relationship.get("to_column", "")







        source_table = source.rsplit(".", 1)[0]



        target_table = target.rsplit(".", 1)[0]







        if (



            source_table in technical_names



            or target_table in technical_names



        ):



            continue







        relationships.append(relationship)







    return relationships











# ============================================================



# FORMATAÇÃO DE TABELAS



# ============================================================







def set_cell_shading(cell, fill: str) -> None:



    properties = cell._tc.get_or_add_tcPr()



    shading = OxmlElement("w:shd")



    shading.set(qn("w:fill"), fill)



    properties.append(shading)











def set_cell_margins(



    cell,



    top: int = 100,



    start: int = 120,



    bottom: int = 100,



    end: int = 120,



) -> None:



    tc_pr = cell._tc.get_or_add_tcPr()



    tc_mar = tc_pr.first_child_found_in("w:tcMar")







    if tc_mar is None:



        tc_mar = OxmlElement("w:tcMar")



        tc_pr.append(tc_mar)







    for margin, value in [



        ("top", top),



        ("start", start),



        ("bottom", bottom),



        ("end", end),



    ]:



        node = tc_mar.find(qn(f"w:{margin}"))







        if node is None:



            node = OxmlElement(f"w:{margin}")



            tc_mar.append(node)







        node.set(qn("w:w"), str(value))



        node.set(qn("w:type"), "dxa")











def set_repeat_table_header(row) -> None:



    tr_pr = row._tr.get_or_add_trPr()



    table_header = OxmlElement("w:tblHeader")



    table_header.set(qn("w:val"), "true")



    tr_pr.append(table_header)











def add_table(document, headers: list, rows: list):



    table = document.add_table(rows=1, cols=len(headers))



    table.style = "Table Grid"



    table.alignment = WD_TABLE_ALIGNMENT.CENTER



    table.autofit = True







    header_row = table.rows[0]



    set_repeat_table_header(header_row)







    for index, header in enumerate(headers):



        cell = header_row.cells[index]



        cell.text = str(header)



        set_cell_shading(cell, DARK_BLUE)



        set_cell_margins(cell)







        for paragraph in cell.paragraphs:



            for run in paragraph.runs:



                run.bold = True



                run.font.name = "Inter"



                run.font.size = Pt(8)



                run.font.color.rgb = RGBColor.from_string(WHITE)







    for row_data in rows:



        cells = table.add_row().cells







        for index, value in enumerate(row_data):



            cells[index].text = "" if value is None else str(value)



            cells[index].vertical_alignment = (



                WD_CELL_VERTICAL_ALIGNMENT.CENTER



            )



            set_cell_margins(cells[index])







            for paragraph in cells[index].paragraphs:



                for run in paragraph.runs:



                    run.font.name = "Inter"



                    run.font.size = Pt(8)



                    run.font.color.rgb = RGBColor.from_string(DARK_GRAY)







    document.add_paragraph()



    return table











# ============================================================



# CONFIGURAÇÃO DO DOCUMENTO



# ============================================================







def configure_document(document) -> None:



    section = document.sections[0]



    section.top_margin = Inches(0.65)



    section.bottom_margin = Inches(0.65)



    section.left_margin = Inches(0.65)



    section.right_margin = Inches(0.65)







    normal = document.styles["Normal"]



    normal.font.name = "Inter"



    normal.font.size = Pt(9)



    normal.font.color.rgb = RGBColor.from_string(DARK_GRAY)



    normal.paragraph_format.space_after = Pt(4)







    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:



        style = document.styles[style_name]



        style.font.name = "Inter"



        style.font.color.rgb = RGBColor.from_string(DARK_BLUE)







    document.styles["Heading 1"].font.size = Pt(20)



    document.styles["Heading 2"].font.size = Pt(14)



    document.styles["Heading 3"].font.size = Pt(10)











def clear_footer(footer) -> None:



    """Remove o conteúdo do rodapé."""







    paragraph = footer.paragraphs[0]



    paragraph.clear()







    for element in list(paragraph._p):



        if element.tag != qn("w:pPr"):



            paragraph._p.remove(element)











def add_page_number(paragraph) -> None:



    """Adiciona o campo PAGE do Word."""







    run = paragraph.add_run()







    begin = OxmlElement("w:fldChar")



    begin.set(qn("w:fldCharType"), "begin")







    instruction = OxmlElement("w:instrText")



    instruction.set(qn("xml:space"), "preserve")



    instruction.text = "PAGE"







    separate = OxmlElement("w:fldChar")



    separate.set(qn("w:fldCharType"), "separate")







    displayed_text = OxmlElement("w:t")



    displayed_text.text = "1"







    end = OxmlElement("w:fldChar")



    end.set(qn("w:fldCharType"), "end")







    run._r.extend([



        begin,



        instruction,



        separate,



        displayed_text,



        end,



    ])











def restart_page_numbering(section, start: int = 1) -> None:



    """Reinicia a numeração da seção."""







    sect_pr = section._sectPr



    page_number_type = sect_pr.find(qn("w:pgNumType"))







    if page_number_type is None:



        page_number_type = OxmlElement("w:pgNumType")



        sect_pr.append(page_number_type)







    page_number_type.set(qn("w:start"), str(start))











def configure_content_footer(section) -> None:



    """Configura o rodapé independente da seção de conteúdo."""







    # Impede que o rodapé seja herdado da seção da capa/sumário.



    section.footer.is_linked_to_previous = False



    section.first_page_footer.is_linked_to_previous = False







    section.different_first_page_header_footer = False







    footer = section.footer



    clear_footer(footer)







    paragraph = footer.paragraphs[0]



    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT







    run = paragraph.add_run("Documentação Técnica  |  ")



    run.font.name = "Inter"



    run.font.size = Pt(8)



    run.font.color.rgb = RGBColor.from_string(DARK_GRAY)







    add_page_number(paragraph)







    restart_page_numbering(section, start=1)











def configure_front_matter_footers(document) -> None:



    """Mantém capa e sumário sem numeração visível."""







    front_section = document.sections[0]



    front_section.different_first_page_header_footer = False







    front_section.footer.is_linked_to_previous = False



    front_section.first_page_footer.is_linked_to_previous = False







    clear_footer(front_section.footer)



    clear_footer(front_section.first_page_footer)











# ============================================================



# CAPA E SUMÁRIO



# ============================================================







def add_cover(document, project_name: str) -> None:



    document.add_paragraph("\n\n\n")







    title = document.add_paragraph()



    title.alignment = WD_ALIGN_PARAGRAPH.CENTER







    run = title.add_run("ANÁLISE DE PROJETO - SCOPE BI")



    run.bold = True



    run.font.name = "Inter"



    run.font.size = Pt(28)



    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)







    subtitle = document.add_paragraph()



    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER







    run = subtitle.add_run("DOCUMENTAÇÃO TÉCNICA")



    run.bold = True



    run.font.name = "Inter"



    run.font.size = Pt(15)



    run.font.color.rgb = RGBColor.from_string(GREEN)







    document.add_paragraph("\n")







    project = document.add_paragraph()



    project.alignment = WD_ALIGN_PARAGRAPH.CENTER







    run = project.add_run(project_name)



    run.bold = True



    run.font.name = "Inter"



    run.font.size = Pt(19)



    run.font.color.rgb = RGBColor.from_string(DARK_GRAY)







    generated = document.add_paragraph()



    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER



    generated.add_run(



        "Documento gerado em "



        + datetime.now().strftime("%d/%m/%Y %H:%M")



    )







    document.add_page_break()















def add_contents_page(document) -> None:



    document.add_heading("Sumário", level=1)



    document.paragraphs[-1].style = "Title"



    paragraph = document.add_paragraph()







    field_begin = OxmlElement("w:fldChar")



    field_begin.set(qn("w:fldCharType"), "begin")







    instruction = OxmlElement("w:instrText")



    instruction.set(qn("xml:space"), "preserve")



    instruction.text = r'TOC \o "1-3" \h \z \u'







    field_separate = OxmlElement("w:fldChar")



    field_separate.set(qn("w:fldCharType"), "separate")







    placeholder = OxmlElement("w:t")



    placeholder.text = (



        "Para atualizar o sumário, clique com o botão direito "



        "nesta área e selecione Atualizar Campo."



    )







    field_end = OxmlElement("w:fldChar")



    field_end.set(qn("w:fldCharType"), "end")







    run = paragraph.add_run()



    run._r.append(field_begin)



    run._r.append(instruction)



    run._r.append(field_separate)



    run._r.append(placeholder)



    run._r.append(field_end)







    paragraph.paragraph_format.space_after = Pt(6)















# ============================================================



# 1. VISÃO GERAL



# ============================================================







def count_report_pages(report_dir: Path | None) -> int | str:
    """Conta as telas usando a ordem oficial do relatório PBIP."""
    if report_dir is None:
        return "-"

    pages_file = Path(report_dir) / "definition" / "pages" / "pages.json"
    if not pages_file.is_file():
        return "-"

    try:
        with pages_file.open("r", encoding="utf-8") as file:
            pages_data = json.load(file)
    except (OSError, json.JSONDecodeError):
        return "-"

    page_order = pages_data.get("pageOrder")
    if isinstance(page_order, list):
        return len(page_order)
    return "-"


def add_overview(document, model: dict, report_dir: Path | None = None) -> None:

    document.add_heading("1. Visão geral", level=1)



    business_tables = get_business_tables(model)

    business_relationships = get_business_relationships(model)

    columns = model.get("columns", [])



    metrics = [

        ("Tabelas", len(business_tables)),

        (

            "Medidas",

            sum(

                1

                for measure in model.get("measures", [])

                if not is_technical_table(measure.get("table", ""))

            ),

        ),

        (

            "Colunas",

            sum(

                1

                for column in columns

                if not is_technical_table(column.get("table", ""))

            ),

        ),

        (

            "Col. Calculadas",

            sum(

                1

                for column in columns

                if not is_technical_table(column.get("table", ""))

                and is_calculated_column(column)

            ),

        ),

        ("Relacionamentos", len(business_relationships)),

        ("Telas", count_report_pages(report_dir)),

    ]



       # Grade com 3 colunas e 2 linhas, mantendo cada KPI independente.

    layout_table = document.add_table(rows=2, cols=3)

    layout_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    layout_table.autofit = False



    card_width = Inches(2.15)



    for row in layout_table.rows:

        for cell in row.cells:

            cell.width = card_width

            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            cell.text = ""



    for index, (label, value) in enumerate(metrics):

        row_index = index // 3

        col_index = index % 3

        container = layout_table.cell(row_index, col_index)



        # Margens da célula externa para separar os cards.

        set_cell_margins(

            container,

            top=70,

            bottom=70,

            start=70,

            end=70,

        )



        card = container.add_table(rows=1, cols=1)

        card.autofit = False

        card.columns[0].width = card_width - Inches(0.15)



        cell = card.cell(0, 0)

        cell.width = card_width - Inches(0.15)

        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        cell.text = ""



        set_cell_shading(cell, LIGHT_GREEN)

        set_cell_margins(

            cell,

            top=240,

            bottom=240,

            start=100,

            end=100,

        )



        # Bordas discretas em cada card.

        tc_pr = cell._tc.get_or_add_tcPr()

        borders = OxmlElement("w:tcBorders")



        for edge in ("top", "left", "bottom", "right"):

            border = OxmlElement(f"w:{edge}")

            border.set(qn("w:val"), "single")

            border.set(qn("w:sz"), "5")

            border.set(qn("w:color"), BORDER_GREEN)

            borders.append(border)



        tc_pr.append(borders)



        paragraph = cell.paragraphs[0]

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        paragraph.paragraph_format.space_after = Pt(0)



        value_run = paragraph.add_run(f"{value}\n")

        value_run.bold = True

        value_run.font.name = "Inter"

        value_run.font.size = Pt(23)

        value_run.font.color.rgb = RGBColor.from_string(DARK_BLUE)



        label_run = paragraph.add_run(label)

        label_run.bold = True

        label_run.font.name = "Inter"

        label_run.font.size = Pt(9)

        label_run.font.color.rgb = RGBColor.from_string(DARK_BLUE)



    document.add_paragraph()









# ============================================================



# 2. RELACIONAMENTOS



# ============================================================







def add_table_list(document, model: dict) -> None:



    document.add_heading("2. Relacionamentos", level=1)



    document.add_heading("2.1 Lista de Tabelas", level=2)







    business_tables = get_business_tables(model)



    columns = model.get("columns", [])



    measures = model.get("measures", [])







    rows = []







    for index, table in enumerate(business_tables, start=1):



        name = table.get("name", "")







        column_count = sum(



            1



            for column in columns



            if column.get("table") == name



        )







        measure_count = sum(



            1



            for measure in measures



            if measure.get("table") == name



        )

        calculated_column_count = sum(

            1

            for column in columns

            if column.get("table") == name and is_calculated_column(column)

        )



        source = "DAX" if name == "d_Calendario" else "Importação"



        rows.append((

            index,

            name,

            column_count,

            measure_count,

            calculated_column_count,

            source,

        ))







    add_table(



        document,



        ["Nº", "Tabela", "Colunas", "Medidas", "Col. Calc.", "Fonte"],



        rows,



    )











def add_relationship_list(document, model: dict) -> None:



    document.add_heading("2.2 Lista de Relacionamentos", level=2)







    relationships = get_business_relationships(model)







    document.add_paragraph(



        f"O modelo possui {len(relationships)} relacionamentos "



        "entre tabelas não técnicas."



    )







    rows = []







    for relationship in relationships:



        source = relationship.get("from_column", "")



        target = relationship.get("to_column", "")







        bidirectional = (



            "Sim"



            if relationship.get("cross_filtering_behavior")



            == "bothDirections"



            else "Não"



        )







        active = (



            "Sim"



            if relationship.get("is_active") is True



            else "Não"



            if relationship.get("is_active") is False



            else "Não informado"



        )







        rows.append((



            source,



            "→",



            target,



            bidirectional,



            active,



        ))







    add_table(



        document,



        ["Origem", "", "Destino", "Bidirecional", "Ativo"],



        rows,



    )











# ============================================================



# 3. LISTA DE TELAS



# ============================================================







def add_report_pages(document) -> None:



    document.add_heading("3. Lista de telas", level=1)











# ============================================================



# 4. ANEXOS DAS MEDIDAS



# ============================================================







def infer_unit(measure: dict) -> str:



    """Infere uma unidade básica a partir do nome da medida."""







    name = measure.get("name", "").casefold()



    expression = measure.get("expression", "").casefold()







    if "faturamento" in name or "receita" in name:



        return "R$"







    if (



        "%" in name



        or "taxa" in name



        or "percentual" in name



    ):



        return "%"







    if "ticket" in name or "preço" in name or "valor" in name:



        return "R$"







    if "count(" in expression or "countrows(" in expression:



        return "Unidades"







    return "—"













def add_measures(document, model: dict) -> None:

    document.add_heading("4. Anexos das medidas", level=1)



    rows = []



    for index, measure in enumerate(model.get("measures", []), start=1):

        name = measure.get("name", "Medida sem nome")



        expression = (

            measure.get("expression", "").strip()

            or "Expressão DAX não disponível."

        )



        rows.append((

            index,

            name.replace("_", " "),

            expression,

            infer_unit(measure),

        ))



    add_table(

        document,

        ["Nº", "Descrição", "Fórmula", "Unidade"],

        rows,

    )











# ============================================================



# GERAÇÃO DO DOCX



# ============================================================







def generate_docx(



    model: dict,



    output_path: Path,



    project_name: str,



    report_dir: Path | None = None,



) -> Path:



    """



    Gera a documentação DOCX com:



    - Capa sem número de página



    - Sumário sem número de página



    - Conteúdo iniciado em uma nova seção, com página 1



    """







    output_path = Path(output_path)



    output_path.parent.mkdir(parents=True, exist_ok=True)







    document = Document()



    configure_document(document)







    # Seção 1: capa e sumário, sem rodapé visível.



    add_cover(document, project_name)



    add_contents_page(document)







    configure_front_matter_footers(document)







    # Cria uma nova seção em nova página após o sumário.



    content_section = document.add_section(WD_SECTION.NEW_PAGE)







    # Mantém margens iguais às da seção anterior.



    content_section.top_margin = Inches(0.65)



    content_section.bottom_margin = Inches(0.65)



    content_section.left_margin = Inches(0.65)



    content_section.right_margin = Inches(0.65)







    # Seção 2: rodapé próprio e numeração reiniciada em 1.



    configure_content_footer(content_section)







    add_overview(document, model, report_dir)



    add_table_list(document, model)



    add_relationship_list(document, model)



    add_report_pages(document)



    add_measures(document, model)







    document.save(output_path)







    if not output_path.is_file() or output_path.stat().st_size == 0:



        raise RuntimeError(



            "O arquivo DOCX não foi gerado corretamente."



        )







    return output_path
