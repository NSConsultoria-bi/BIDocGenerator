
from pathlib import Path
from docx import Document

from tmdl_extractor import extract_tmdl_model
from report_generator import generate_docx


PROJECT_DIR = Path(
    r"C:\Users\Ewerton\OneDrive\Documentos\MBA EXPERIUM\DATA CHALLENGES\1\CASE_OLIST_PBIP"
)

SEMANTIC_MODEL_DIR = PROJECT_DIR / "CASE OLIST_PBIP.SemanticModel"
REPORT_DIR = PROJECT_DIR / "CASE OLIST_PBIP.Report"


def main():
    model = extract_tmdl_model(SEMANTIC_MODEL_DIR)

    output_path = Path(__file__).parent / "documentacao_olist_test.docx"

    generate_docx(
        model=model,
        output_path=output_path,
        project_name="CASE OLIST",
        report_dir=REPORT_DIR,
    )

    document = Document(output_path)

    content = "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )

    tables_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )

    full_text = content + "\n" + tables_text

    required_sections = [
        "E-BI DOC GENERATOR",
        "Visão Geral",
        "Esquema de Relacionamento",
        "Telas do projeto",
        "Anexo de Medidas",
        "Faturamento_Total",
    ]

    print("=== VALIDAÇÃO DO CONTEÚDO DOCX ===")

    for section in required_sections:
        found = section.casefold() in full_text.casefold()
        print(f"{section}: {'OK' if found else 'AUSENTE'}")

    print(f"Parágrafos: {len(document.paragraphs)}")
    print(f"Tabelas do documento: {len(document.tables)}")
    print(f"Arquivo: {output_path}")


if __name__ == "__main__":
    main()
